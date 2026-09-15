"""推理验收：原图区域测量、无掩罩叠加/差异图与真实数组独立复算。"""

import argparse
import csv
import json
from pathlib import Path

import numpy as np
import torch
from openpyxl import load_workbook
from PIL import Image, ImageChops, ImageEnhance

REFERENCE_BOUNDS = {
    "checkpoints": (210, 117, 341, 271),
    "samples": (567, 117, 337, 271),
    "fields": (919, 117, 352, 271),
    "metrics": (1287, 117, 355, 271),
    "settings": (210, 397, 1432, 89),
    "table": (210, 495, 1432, 218),
    "chart": (210, 724, 1432, 194),
}


def visual(root: Path, reference: Path) -> dict:
    """保留全部像素差异，不用遮罩或百分比代替区域测量。"""
    original = Image.open(reference).convert("RGB")
    actual = Image.open(root / "ui/1672x941.png").convert("RGB")
    if original.size != (1672, 941) or actual.size != original.size:
        raise ValueError("原图与实际截图尺寸不一致")
    Image.blend(original, actual, 0.5).save(root / "ui/overlay.png")
    ImageEnhance.Contrast(ImageChops.difference(original, actual)).enhance(3).save(
        root / "ui/difference.png"
    )
    together = Image.new("RGB", (3344, 941), "white")
    together.paste(original, (0, 0))
    together.paste(actual, (1672, 0))
    together.save(root / "ui/side-by-side.png")
    measured = json.loads((root / "ui/1672x941-regions.json").read_text())
    result = []
    for region, expected in REFERENCE_BOUNDS.items():
        errors = {
            axis: abs(measured[region][axis] - v)
            for axis, v in zip(("x", "y", "width", "height"), expected, strict=True)
        }
        result.append(
            {
                "region": region,
                "expected": expected,
                "actual": measured[region],
                "errors": errors,
                "passed": max(errors.values()) <= 4,
            }
        )
    report = {
        "regions": result,
        "passed": all(r["passed"] for r in result),
        "differences": "平台品牌与任务元数据沿用现行产品；字段/指标数量来自实际目录；分页50；图表使用样本等权统计，非原图示意数值；字体抗锯齿不同。",
    }
    (root / "ui/measurements.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    if not report["passed"]:
        raise AssertionError(report)
    return report


def numerical(root: Path, browser: str = "real-browser") -> dict:
    """直接读取预测和真值独立复算，不调用产品评价函数。"""
    evidence = json.loads((root / browser / "results.json").read_text())
    context = json.loads((root / "real-cfd/context.json").read_text())
    project = Path(context["project_directory"])
    rows = []
    for item in evidence["results"]["items"]:
        ref = next(f for f in item["files"] if f["name"] == "manifest.json")
        path = project / ref["path"]
        manifest = json.loads(path.read_text())
        # Server project根相对路径。版本引用只定位本次已提交manifest。
        for record in item["metric_records"]:
            key = manifest["domains"][record["domain"]]["targets"][record["field"]]
            p = (
                torch.load(
                    path.parent / manifest["filemap"][key + ".prediction"], weights_only=True
                )
                .numpy()
                .astype(np.float64)
            )
            t = (
                torch.load(path.parent / manifest["filemap"][key + ".truth"], weights_only=True)
                .numpy()
                .astype(np.float64)
            )
            component = record["component"]
            if component == "magnitude":
                p, t = np.linalg.norm(p, axis=1), np.linalg.norm(t, axis=1)
            else:
                p, t = (
                    p[:, 0 if component == "scalar" else int(component)],
                    t[:, 0 if component == "scalar" else int(component)],
                )
            diff = p - t
            calculated = {
                "relative_l2": float(np.linalg.norm(diff) / np.linalg.norm(t))
                if np.linalg.norm(t)
                else None,
                "mae": float(np.abs(diff).mean()),
                "rmse": float(np.sqrt(np.mean(diff**2))),
                "max_abs_error": float(np.max(np.abs(diff))),
                "r2": float(1 - np.sum(diff**2) / np.sum((t - t.mean()) ** 2))
                if np.sum((t - t.mean()) ** 2)
                else None,
            }
            for name, actual in record["values"].items():
                expected = calculated[name]
                if expected is None:
                    assert actual is None
                else:
                    assert np.isclose(actual, expected, rtol=1e-12, atol=1e-12), (
                        name,
                        actual,
                        expected,
                    )
            rows.append(
                {
                    "run_id": item["run_id"],
                    "split": item["split"],
                    "sample": item["sample"],
                    "points": len(p),
                    "metrics": calculated,
                }
            )
    with (root / browser / "results.csv").open(encoding="utf-8-sig") as stream:
        exported = list(csv.DictReader(stream))
    workbook = load_workbook(root / browser / "results.xlsx", read_only=True, data_only=True)
    assert {"统计", "逐样本指标", "来源", "指标说明"} <= set(workbook.sheetnames)
    assert workbook["统计"].max_row == len(exported) + 1
    expected = {
        r["checkpoint_id"]: r
        for r in evidence["results"]["statistics"]
        if r["field_id"] == exported[0]["field_id"]
        and r["metric"] == exported[0]["metric"]
        and r["split"] == exported[0]["split"]
    }
    for row in exported:
        assert float(row["mean"]) == expected[row["checkpoint_id"]]["mean"]
    report = {
        "rows": rows,
        "csv_rows": len(exported),
        "xlsx_sheets": workbook.sheetnames,
        "independent_metrics": "passed",
        "production_accuracy_claim": False,
    }
    (root / browser / "independent-check.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2)
    )
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument(
        "--reference", type=Path, default=Path("docs/prototypes/dojo-inference-reference.png")
    )
    parser.add_argument("--visual-only", action="store_true")
    parser.add_argument("--browser-evidence", default="real-browser")
    args = parser.parse_args()
    visual(args.root, args.reference)
    if not args.visual_only:
        numerical(args.root, args.browser_evidence)
