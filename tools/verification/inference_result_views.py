"""独立复核实际浏览器下载的两种推理表格，不调用业务导出实现。"""

import argparse
import csv
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


def verify(root: Path) -> dict:
    """按固定逐样本值重建数值预期，并核对CSV、Excel和页面列顺序。"""
    evidence = json.loads((root / "real-browser-evidence.json").read_text())
    records = evidence["results"]["records"]
    counts = {}
    for mode in ("checkpoint", "sample"):
        export = next(r for r in evidence["exports"] if r["selection"]["view"]["mode"] == mode)
        config = export["selection"]["view"]
        with (root / f"real-{mode}.csv").open(encoding="utf-8-sig") as stream:
            csv_rows = list(csv.reader(stream))
        book = load_workbook(root / f"real-{mode}.xlsx", data_only=True)
        assert book.sheetnames[0] == "当前表格"
        assert {"统计", "逐样本指标", "指标说明", "来源"} <= set(book.sheetnames)
        xlsx_rows = list(book["当前表格"].values)
        assert csv_rows[0] == list(xlsx_rows[0]) == evidence[mode]["headers"]
        assert len(csv_rows) == len(xlsx_rows) == len(evidence[mode]["rows"]) + 1
        for index, (csv_row, xlsx_row) in enumerate(zip(csv_rows[1:], xlsx_rows[1:], strict=True)):
            assert csv_row[:2] == [str(v) for v in xlsx_row[:2]]
            assert csv_row[:2] == evidence[mode]["rows"][index][:2]
            if mode == "checkpoint":
                subset = [r for r in records if r["checkpoint"] == xlsx_row[0]]
                assert xlsx_row[1] == len({(r["split"], r["sample"]) for r in subset})
            else:
                split = {"训练集": "train", "验证集": "eval", "测试集": "test"}[xlsx_row[1]]
                subset = [
                    r
                    for r in records
                    if r["checkpoint_id"] == config["checkpoint_id"]
                    and r["sample"] == xlsx_row[0]
                    and r["split"] in (split, "validation" if split == "eval" else split)
                ]
            expected = []
            for pair in config["pairs"]:
                values = np.asarray(
                    [
                        r["values"][pair["metric"]]
                        for r in subset
                        if r["field_id"] == pair["field_id"]
                    ],
                    dtype=np.float64,
                )
                assert len(values)
                if mode == "sample":
                    assert len(values) == 1
                    expected.append(values[0])
                else:
                    for aggregate in config["aggregations"]:
                        expected.append(
                            {
                                "mean": np.mean(values),
                                "median": np.median(values),
                                "p90": np.quantile(values, 0.9, method="linear"),
                                "max": np.max(values),
                            }[aggregate]
                        )
            np.testing.assert_allclose(
                [float(v) for v in csv_row[2:]], expected, rtol=1e-12, atol=1e-12
            )
            np.testing.assert_allclose(xlsx_row[2:], expected, rtol=1e-12, atol=1e-12)
            np.testing.assert_allclose(
                [float(v) for v in evidence[mode]["rows"][index][2:]],
                expected,
                rtol=5e-5,
                atol=1e-10,
            )
        counts[mode] = {
            "rows": len(csv_rows) - 1,
            "columns": len(csv_rows[0]),
            "numeric_cells": (len(csv_rows) - 1) * (len(csv_rows[0]) - 2),
        }
    assert (
        evidence["predictions"] == 0 and not evidence["errors"] and not evidence["failedRequests"]
    )
    return {
        "status": "passed",
        "batch": evidence["batch"],
        "fixed_records": len(records),
        "downloads": counts,
        "prediction_submissions": 0,
        "method": "numpy float64 from fixed sample metrics; P90 linear; XLSX/CSV readback",
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("root", type=Path)
    args = parser.parse_args()
    result = verify(args.root)
    (args.root / "download-verification.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2)
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
