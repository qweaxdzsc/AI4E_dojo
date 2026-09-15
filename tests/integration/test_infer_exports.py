"""真实CSV和XLSX读回，不以文件扩展名冒充表格格式。"""

import csv

import pytest

from openpyxl import load_workbook

from ai4e_core.abilities.report import export_tables


def test_csv_xlsx_numeric_and_source(tmp_path):
    tables = {
        "统计": [{"checkpoint": "=unsafe", "mean": 0.125, "undefined": None, "unit": "Pa"}],
        "来源": [{"split": "validation", "sample": "中文样本", "algorithm": "physical-metrics-v2"}],
    }
    csv_path = tmp_path / "metrics.csv"
    xlsx_path = tmp_path / "metrics.xlsx"
    export_tables(csv_path, tables, format="csv")
    export_tables(xlsx_path, tables, format="xlsx")
    with csv_path.open(encoding="utf-8-sig") as stream:
        rows = list(csv.DictReader(stream))
    assert rows[0]["mean"] == "0.125" and rows[0]["checkpoint"] == "'=unsafe"
    book = load_workbook(xlsx_path)
    assert book.sheetnames == ["统计", "来源"]
    assert book["统计"]["B2"].value == 0.125 and book["统计"]["B2"].data_type == "n"
    assert book["统计"]["A2"].data_type == "s" and book["来源"]["B2"].value == "中文样本"


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_fixed_post_xlsx_preserves_values_and_identity(tmp_path, format):
    from ai4e_core.applications.aero_cfd.post.result_evaluation import export_evaluation

    path = tmp_path / ("post." + format)
    export_evaluation({"rows": [{"id": "r", "split": "validation", "sample": "car",
        "algorithm": "physical-metrics-v2", "unit": "Pa", "values": {"mae": 0.125}}]},
        path, format=format)
    if format == "xlsx":
        rows = list(load_workbook(path).active.values)
        value = dict(zip(rows[0], rows[1]))
    else:
        with path.open(encoding="utf-8-sig") as stream:
            value = next(csv.DictReader(stream))
        value["mae"] = float(value["mae"])
    assert value["mae"] == 0.125 and value["split"] == "validation"
    assert value["algorithm"] == "physical-metrics-v2" and value["unit"] == "Pa"
