"""显式表格导出；数值保持数值，文本不作为电子表格公式执行。"""

import csv
import json
from pathlib import Path

from ai4e_core.abilities.data.save.arrays import atomic_path


def _value(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_tables(path, tables: dict[str, list[dict]], *, format: str) -> str:
    """CSV消费首张表，XLSX按输入顺序建立各表；调用方决定稳定来源。"""
    if format not in {"csv", "xlsx"} or not tables:
        raise ValueError("导出需要表格及csv/xlsx格式")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with atomic_path(path) as temporary:
        if format == "csv":
            rows = next(iter(tables.values()))
            names = list(dict.fromkeys(k for row in rows for k in row))
            with temporary.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=names)
                writer.writeheader()
                writer.writerows({k: _value(v) for k, v in r.items()} for r in rows)
        else:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill

            book = Workbook()
            book.remove(book.active)
            for name, rows in tables.items():
                sheet = book.create_sheet(name[:31])
                names = list(dict.fromkeys(k for row in rows for k in row))
                sheet.append(names)
                for row in rows:
                    sheet.append([_value(row.get(k)) for k in names])
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="173A78")
                    cell.fill = PatternFill("solid", fgColor="EEF5FF")
                for column in sheet.columns:
                    sheet.column_dimensions[column[0].column_letter].width = min(
                        50, max(15, max(len(str(c.value or "")) for c in column) + 2)
                    )
            book.save(temporary)
    return str(path)
